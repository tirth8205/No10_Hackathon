"""Build the realistic demo documents: letterheaded Word files, PDFs, phone photos and a police email.

    .venv/bin/python demo/make_documents.py

Writes demo/files/<case-slug>/<stem>.{docx,pdf,jpg} for every demo case in demo/seed_demo.py, a logo PNG per
organisation under demo/files/logos/, the police email as a real .eml, and the corrected Case A letter again at
demo/files/case-a-support-letter-corrected.{pdf,jpg} for the live re-upload.

The letter wording is the case text from seed_demo, sentence for sentence, so the known verdicts still hold.
Only a letterhead (logo, organisation, address, phone, email, website), a date and salutation where the text has
none, a signature block and a footer are added. Where the source text ends with its own name / role /
organisation lines, those lines become the signature block and the organisation line moves to the letterhead.

Everything is fictional: organisations, people, addresses, Ofcom drama phone numbers (0113 496 0xxx and so on),
unregistered domains and invented charity numbers. The one exception is the Women's Aid letter, worded like
MoJ's own sample, which carries the real www.womensaid.org.uk because the verification check fetches it.

Deterministic: fixed seed for the photo noise, fixed Word metadata and zip timestamps, and the PDFs are
rewritten with a fixed creation date and document id. LibreOffice (soffice) renders the PDF and the page image.
"""
from __future__ import annotations

import email.message
import email.utils
import hashlib
import io
import random
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from pathlib import Path

from PIL import Image, ImageChops, ImageDraw, ImageFilter, ImageFont

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from demo import seed_demo as S  # noqa: E402  (the case texts)

OUT = ROOT / "demo" / "files"
LOGOS = OUT / "logos"
SOFFICE = shutil.which("soffice") or "/opt/homebrew/bin/soffice"
LO_PROFILE = Path(tempfile.gettempdir()) / "gateway-demo-lo-profile"  # own profile: works while LibreOffice is open
SEED = 20261005
PHOTO_WIDTH = 1600
PNG_EXPORT = 'png:writer_png_Export:{"PixelWidth":{"type":"long","value":"2480"},"PixelHeight":{"type":"long","value":"3508"}}'
UK = timezone(timedelta(hours=1))
GREY = (0x59, 0x59, 0x59)
FONT_CANDIDATES = ["/System/Library/Fonts/Supplemental/Arial Bold.ttf", "/Library/Fonts/Arial Bold.ttf",
                   "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"]


# --------------------------------------------------------------------------- organisations (fictional)
@dataclass(frozen=True)
class Org:
    key: str
    name: str
    initials: str
    colour: tuple[int, int, int]
    address: str | None = None  # "\n" allowed for a multi-line address
    phone: str | None = None
    email: str | None = None
    website: str | None = None
    footer: str | None = None
    logo: bool = True


ORGS = {o.key: o for o in [
    Org("brookfield", "Brookfield Women's Support Service", "BW", (118, 42, 110),
        "PO Box 318, Leeds LS2 7EQ", "0113 496 0312", "info@brookfieldwss.org.uk", None,
        "Brookfield Women's Support Service is a registered charity in England and Wales, no. 1173209."),
    Org("moorside", "Moorside Support Service", "MS", (0, 96, 124),
        "PO Box 77, Sheffield S1 2BX", "0114 496 0455", "info@moorsidesupport.org.uk", None,
        "Moorside Support Service is a registered charity, no. 1162844, and a company limited by guarantee, no. 09418217."),
    Org("elmgrove", "Elm Grove Medical Practice", "EG", (0, 114, 82),
        "8 Elm Grove\nLeeds\nLS1 4AP", "0113 496 0148", "elmgrove.reception@nhs.net", None,
        "Partners: Dr A Rowe, Dr M Kaur, Dr S Oyelaran.  Out of hours: call 111."),
    Org("wrenfield", "Wrenfield House Refuge", "WH", (30, 86, 152),
        "Correspondence: PO Box 1402, Bristol BS1 6PN", "0117 496 0233", "admin@wrenfieldhouse.org.uk", None,
        "Wrenfield House Refuge is a registered charity in England and Wales, no. 1156317. The refuge address is confidential."),
    Org("larkspur", "Larkspur Women's Centre", "LW", (166, 62, 32),
        "PO Box 51, Nottingham NG1 5DT", "0115 496 0170", "referrals@larkspurwc.org.uk", "www.larkspurwc.org.uk",
        "Larkspur Women's Centre is a registered charity in England and Wales, no. 1128406."),
    Org("womensaid", "Women's Aid", "WA", (60, 60, 60), footer="www.womensaid.org.uk", logo=False),  # plain, as in MoJ's sample
    Org("newpathways", "New Pathways Support Trust", "NP", (92, 112, 32),
        "PO Box 1180, Cardiff CF10 1EP", "029 2018 0421", "office@newpathwayssupport.org.uk", None,
        "New Pathways Support Trust is a registered charity in England and Wales, no. 1151920."),
    Org("northfield", "Northfield Surgery", "NS", (0, 92, 142),
        "14 Station Road\nSheffield\nS1 2HH", "0114 496 0000", None, None,
        "Northfield Surgery is an NHS GP practice.  Out of hours: call 111."),
]}


# --------------------------------------------------------------------------- letters
@dataclass
class Letter:
    case: str                      # case slug = folder under demo/files
    stem: str                      # file name without extension
    org: str | None                # ORGS key; None for a plain statement
    text: str                      # the case text, kept sentence for sentence
    date: str | None = None        # printed under the letterhead; None when the text carries its own date
    salutation: bool = False       # add "To whom it may concern" (only when the text has none)
    strip_leading: int = 0         # letterhead lines already in the text, rendered by the header instead
    strip_trailing: int = 0        # name / role / organisation lines at the end, rendered as the signature block
    signature: list[str] = field(default_factory=list)
    trailer: str | None = None     # small print after the signature (Case D's prompt injection)
    statement: bool = False        # Leah Clarke: a statement, not a letter
    pdf: bool = True
    photo: bool = True
    file_date: str | None = None   # for the Word / PDF metadata when `date` is None

    @property
    def when(self) -> datetime:
        return datetime.strptime(self.date or self.file_date, "%d %B %Y").replace(hour=9, minute=30, tzinfo=UK)


def _case(name: str) -> dict:
    return next(s for s in S.CASES if s["applicant_name"] == name)


def _text(name: str, i: int = 0) -> str:
    return _case(name)["docs"][i][1]


def letters() -> list[Letter]:
    amira = _text("Amira Hussain")
    assert amira.endswith(S.INJECTION)
    return [
        Letter("jane-doe", "support-letter", "brookfield", _text("Jane Doe"), "29 September 2026", True,
               signature=["Sam Ellis", "Senior Support Worker"]),
        Letter("jane-doe", "support-letter-corrected", "brookfield", S.CORRECTED_A, "2 October 2026", True,
               signature=["Sam Ellis", "Senior Support Worker"]),
        Letter("joe-bloggs", "support-letter", "moorside", _text("Joe Bloggs"), "25 September 2026", True,
               signature=["Priya Nair", "Senior Support Worker"]),
        Letter("jane-smith", "gp-letter", "elmgrove", _text("Jane Smith", 0), "22 September 2026", True,
               signature=["Dr Anita Rowe", "General Practitioner", "GMC number 7593844"]),
        Letter("jane-smith", "refuge-letter", "wrenfield", _text("Jane Smith", 1), "30 September 2026",
               strip_trailing=3, signature=["John Smith", "Domestic Abuse Support Officer"]),
        Letter("amira-hussain", "support-letter", "larkspur", amira[: -len(S.INJECTION)], "26 September 2026", True,
               strip_trailing=3, signature=["Kerry Hussain", "Senior Outreach Worker"], trailer=S.INJECTION.strip()),
        Letter("natalia-hussain", "refuge-letter", "womensaid", _text("Natalia Hussain"), "21 August 2026",
               strip_trailing=3, signature=["Mohammed Walsh", "Domestic Abuse Support Officer"]),
        Letter("hannah-shah", "refuge-letter", "newpathways", _text("Hannah Shah"), "17 September 2026",
               strip_trailing=3, signature=["Fatima Begum", "Domestic Abuse Support Officer"]),
        Letter("leah-clarke", "statement", None, _text("Leah Clarke"), statement=True, pdf=False, photo=False,
               signature=["Signed: Leah Clarke", "Date: 28 September 2026"], file_date="28 September 2026"),
        Letter("lucy-brown", "gp-letter", "northfield", S.GP_LETTER_WITH_LETTERHEAD, strip_leading=5, strip_trailing=4,
               signature=["Dr Rachel Hughes", "General Practitioner", "Northfield Surgery", "r.hughes@nhs.net"],
               file_date="30 September 2026"),
    ]


EMAIL = ("sarah-mensah", "police-email.eml", S.POLICE_EMAIL, datetime(2026, 9, 4, 14, 12, 37, tzinfo=UK))


def body_blocks(letter: Letter) -> list[str]:
    """Paragraph blocks of the body after lifting out the letterhead and signature lines."""
    lines = letter.text.strip().splitlines()
    for _ in range(letter.strip_leading):
        while lines and not lines[0].strip():
            lines.pop(0)
        lines.pop(0)
    for _ in range(letter.strip_trailing):
        while lines and not lines[-1].strip():
            lines.pop()
        lines.pop()
    return [b.strip() for b in re.split(r"\n\s*\n", "\n".join(lines).strip()) if b.strip()]


# --------------------------------------------------------------------------- logo
def _font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for f in FONT_CANDIDATES:
        if Path(f).exists():
            return ImageFont.truetype(f, size)
    return ImageFont.load_default(size=size)


def make_logo(org: Org, path: Path) -> None:
    size = 256
    img = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((6, 6, size - 6, size - 6), radius=58, fill=org.colour + (255,))
    light = tuple(min(255, c + 40) for c in org.colour) + (255,)
    d.rounded_rectangle((6, 6, size - 6, size // 2), radius=58, fill=light)
    d.rectangle((6, size // 2 - 60, size - 6, size // 2), fill=light)  # square off the band's lower corners
    d.text((size / 2, size / 2 + 4), org.initials, font=_font(112), fill=(255, 255, 255, 255), anchor="mm")
    img.save(path, "PNG", optimize=True)


# --------------------------------------------------------------------------- Word
def _rule_below(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "4"), ("w:color", "A6A6A6")):
        bottom.set(qn(k), v)
    pbdr.append(bottom)
    pPr.insert_element_before(pbdr, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                              "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
                              "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                              "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
                              "w:rPr", "w:sectPr", "w:pPrChange")


def _para(doc, text: str = "", *, size: float | None = None, bold: bool = False, colour: tuple | None = None,
          before: float | None = None, after: float | None = None, align=None):
    from docx.shared import Pt, RGBColor

    p = doc.add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if colour:
            r.font.color.rgb = RGBColor(*colour)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    return p


def build_docx(letter: Letter, logo: Path | None, out: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    org = ORGS[letter.org] if letter.org else None
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(22)
    sec.top_margin, sec.bottom_margin = Mm(18), Mm(18)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.12

    if org:
        if org.logo and logo:
            p = _para(doc, after=3)
            p.add_run().add_picture(str(logo), width=Mm(15))
        _para(doc, org.name, size=15, bold=True, colour=org.colour, after=1)
        lines = []
        if org.address:
            lines.append(org.address)
        contact = "   ·   ".join(x for x in [f"Tel: {org.phone}" if org.phone else None, org.email, org.website] if x)
        if contact:
            lines.append(contact)
        last = doc.paragraphs[-1]
        for ln in lines:
            last = _para(doc, ln, size=8.5, colour=GREY, after=0)
        _rule_below(last)
        last.paragraph_format.space_after = Pt(16)

    if letter.date:
        _para(doc, letter.date, after=12)
    if letter.salutation:
        _para(doc, "To whom it may concern", after=10)

    blocks = body_blocks(letter)
    for i, block in enumerate(blocks):
        if letter.statement and i == 0:
            title, _, rest = block.partition("\n")
            _para(doc, title, size=14, bold=True, after=4)
            if rest:
                _para(doc, rest, after=12)
        else:
            _para(doc, block)

    if letter.signature:
        p = _para(doc, letter.signature[0], bold=True, before=22, after=0)
        for ln in letter.signature[1:]:
            _para(doc, ln, after=0)
    if letter.trailer:
        _para(doc, letter.trailer, size=7, colour=(0xB7, 0xB7, 0xB7), before=28)

    if org and org.footer:
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run(org.footer)
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor(*GREY)

    cp = doc.core_properties
    cp.author = cp.last_modified_by = letter.signature[0].replace("Signed: ", "") if letter.signature else (org.name if org else "")
    cp.created = cp.modified = letter.when.astimezone(timezone.utc).replace(tzinfo=None)
    cp.title = f"{org.name} letter" if org else blocks[0]
    cp.revision = 1
    doc.save(out)
    _stable_zip(out)


def _stable_zip(path: Path) -> None:
    """Rewrite the .docx zip with fixed entry timestamps so the file is byte-identical on every build."""
    src = zipfile.ZipFile(io.BytesIO(path.read_bytes()))
    names = sorted(src.namelist(), key=lambda n: (n != "[Content_Types].xml", n))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in names:
            zi = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            dst.writestr(zi, src.read(name))
    path.write_bytes(buf.getvalue())


# --------------------------------------------------------------------------- LibreOffice
def soffice(convert_to: str, outdir: Path, inputs: list[Path]) -> None:
    if not inputs:
        return
    cmd = [SOFFICE, f"-env:UserInstallation={LO_PROFILE.as_uri()}", "--headless", "--norestore", "--nologo",
           "--convert-to", convert_to, "--outdir", str(outdir), *map(str, inputs)]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=900)
    if r.returncode != 0:
        raise RuntimeError(f"soffice failed ({r.returncode}): {r.stderr.strip() or r.stdout.strip()}")


def _pdf_date(dt: datetime) -> str:
    return dt.astimezone(timezone.utc).strftime("D:%Y%m%d%H%M%SZ")


def stabilise_pdf(path: Path, letter: Letter, docx_bytes: bytes) -> None:
    """LibreOffice stamps the current time and a fresh id into every PDF; fix both so rebuilds are byte-identical."""
    from pypdf import PdfWriter
    from pypdf.generic import ArrayObject, ByteStringObject

    w = PdfWriter(clone_from=str(path))
    w.add_metadata({"/CreationDate": _pdf_date(letter.when), "/ModDate": _pdf_date(letter.when),
                    "/Producer": "LibreOffice Writer", "/Creator": "Writer", "/Author": letter.signature[0] if letter.signature else ""})
    digest = hashlib.sha256(docx_bytes).digest()[:16]
    w._ID = ArrayObject([ByteStringObject(digest), ByteStringObject(digest)])
    xmp = w._root_object.get("/Metadata")  # the XMP packet carries the build time too
    if xmp is not None:
        stream = xmp.get_object()
        stamp = letter.when.strftime("%Y-%m-%dT%H:%M:%S+01:00").encode()
        stream.set_data(re.sub(rb"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}[+-]\d{2}:\d{2}", stamp, stream.get_data()))
    buf = io.BytesIO()
    w.write(buf)
    path.write_bytes(buf.getvalue())


# --------------------------------------------------------------------------- phone photo
def phone_photo(page_png: Path, out: Path, seed: int) -> None:
    rng = random.Random(seed)
    page = Image.open(page_png).convert("RGB")
    w = 1700
    page = page.resize((w, round(page.height * w / page.width)), Image.LANCZOS)
    page = ImageChops.multiply(page, Image.new("RGB", page.size, (251, 247, 238)))  # warm indoor light on the paper
    desk = (66, 58, 54)
    tilted = page.rotate(1.2, resample=Image.BICUBIC, expand=True, fillcolor=desk)
    canvas = Image.new("RGB", (tilted.width + 90, tilted.height + 90), desk)
    canvas.paste(tilted, (45, 45))
    vign = Image.radial_gradient("L").resize(canvas.size, Image.BILINEAR).point(lambda v: 255 - int(v * 0.42))
    canvas = ImageChops.multiply(canvas, Image.merge("RGB", (vign, vign, vign)))
    noise = Image.frombytes("L", canvas.size, rng.randbytes(canvas.width * canvas.height))
    canvas = Image.blend(canvas, Image.merge("RGB", (noise, noise, noise)), 0.07)
    canvas = canvas.filter(ImageFilter.GaussianBlur(0.55))
    canvas = canvas.resize((PHOTO_WIDTH, round(canvas.height * PHOTO_WIDTH / canvas.width)), Image.LANCZOS)
    canvas.save(out, "JPEG", quality=70, optimize=True)  # no EXIF, no timestamp


# --------------------------------------------------------------------------- email
def make_eml(text: str, when: datetime, out: Path) -> None:
    head, body = text.split("\n\n", 1)
    hdrs = dict(line.split(": ", 1) for line in head.splitlines())
    msg = email.message.EmailMessage()
    msg["From"] = hdrs["From"]
    msg["To"] = hdrs["To"]
    msg["Date"] = email.utils.format_datetime(when)
    msg["Subject"] = hdrs["Subject"]
    sender = email.utils.parseaddr(hdrs["From"])[1]
    msg["Message-ID"] = f"<{when.strftime('%Y%m%d%H%M%S')}.{sender.split('@')[0]}@{sender.split('@')[1]}>"
    msg.set_content(body.strip() + "\n")
    out.write_bytes(msg.as_bytes())


# --------------------------------------------------------------------------- build
def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", s.replace("’", "'")).strip()


def build() -> list[Path]:
    OUT.mkdir(exist_ok=True)
    LOGOS.mkdir(exist_ok=True)
    written: list[Path] = []
    for org in ORGS.values():
        if org.logo:
            p = LOGOS / f"{org.key}.png"
            make_logo(org, p)
            written.append(p)

    todo = letters()
    with tempfile.TemporaryDirectory() as tmp:
        stage = Path(tmp)
        staged: dict[str, tuple[Letter, Path]] = {}
        for lt in todo:
            folder = OUT / lt.case
            folder.mkdir(exist_ok=True)
            logo = LOGOS / f"{lt.org}.png" if lt.org and ORGS[lt.org].logo else None
            docx_path = folder / f"{lt.stem}.docx"
            build_docx(lt, logo, docx_path)
            written.append(docx_path)
            key = f"{lt.case}--{lt.stem}"
            shutil.copyfile(docx_path, stage / f"{key}.docx")
            staged[key] = (lt, docx_path)
        pdf_dir, png_dir = stage / "pdf", stage / "png"
        pdf_dir.mkdir()
        png_dir.mkdir()
        soffice("pdf", pdf_dir, [stage / f"{k}.docx" for k, (lt, _) in staged.items() if lt.pdf])
        soffice(PNG_EXPORT, png_dir, [stage / f"{k}.docx" for k, (lt, _) in staged.items() if lt.photo])
        for key, (lt, docx_path) in staged.items():
            if lt.pdf:
                pdf = OUT / lt.case / f"{lt.stem}.pdf"
                shutil.move(pdf_dir / f"{key}.pdf", pdf)
                stabilise_pdf(pdf, lt, docx_path.read_bytes())
                written.append(pdf)
            if lt.photo:
                jpg = OUT / lt.case / f"{lt.stem}.jpg"
                phone_photo(png_dir / f"{key}.png", jpg, SEED + sum(map(ord, key)))
                written.append(jpg)

    case, fn, text, when = EMAIL
    (OUT / case).mkdir(exist_ok=True)
    make_eml(text, when, OUT / case / fn)
    written.append(OUT / case / fn)

    for ext in ("pdf", "jpg"):  # the file dropped live in the demo
        dst = OUT / f"case-a-support-letter-corrected.{ext}"
        shutil.copyfile(OUT / "jane-doe" / f"support-letter-corrected.{ext}", dst)
        written.append(dst)
    return written


def verify(written: list[Path]) -> list[str]:
    from pypdf import PdfReader

    from engine.ingest import ingest

    problems: list[str] = []
    for lt in letters():
        expected = body_blocks(lt) + lt.signature + ([lt.trailer] if lt.trailer else [])
        if lt.org:
            expected.append(ORGS[lt.org].name)
        if lt.pdf:
            pdf = OUT / lt.case / f"{lt.stem}.pdf"
            text = _norm("\n".join(p.extract_text() or "" for p in PdfReader(str(pdf)).pages))
            for block in expected:
                if _norm(block) not in text:
                    problems.append(f"{pdf.relative_to(ROOT)}: text layer is missing {block[:60]!r}")
            ing = ingest(pdf.read_bytes(), pdf.name)
            if ing.kind != "pdf" or ing.needs_transcription or ing.pages != 1:
                problems.append(f"{pdf.relative_to(ROOT)}: ingest gave kind={ing.kind} pages={ing.pages} needs_transcription={ing.needs_transcription}")
        docx_path = OUT / lt.case / f"{lt.stem}.docx"
        ing = ingest(docx_path.read_bytes(), docx_path.name)
        if ing.kind != "docx":
            problems.append(f"{docx_path.relative_to(ROOT)}: ingest kind {ing.kind}")
        dtext = _norm(ing.text)
        for block in expected:
            if _norm(block) not in dtext and block != (ORGS[lt.org].footer if lt.org else None):
                problems.append(f"{docx_path.relative_to(ROOT)}: docx text is missing {block[:60]!r}")
        if lt.photo:
            jpg = OUT / lt.case / f"{lt.stem}.jpg"
            im = Image.open(jpg)
            if im.width != PHOTO_WIDTH:
                problems.append(f"{jpg.relative_to(ROOT)}: width {im.width}")
            if im.getexif():
                problems.append(f"{jpg.relative_to(ROOT)}: carries EXIF")
            ing = ingest(jpg.read_bytes(), jpg.name)
            if ing.kind != "image" or not ing.needs_transcription:
                problems.append(f"{jpg.relative_to(ROOT)}: ingest kind {ing.kind}")
    case, fn, text, _ = EMAIL
    eml = OUT / case / fn
    ing = ingest(eml.read_bytes(), eml.name)
    key = _norm(text.split("\n\n", 1)[1].split("\n\n")[1])
    if ing.kind != "email" or key not in _norm(ing.text) or "pnn.police.uk" not in ing.text:
        problems.append(f"{eml.relative_to(ROOT)}: ingest kind {ing.kind} or body missing")
    for spec in S.CASES:
        for fn, rel in spec.get("files", []):
            if not (OUT / rel).exists():
                problems.append(f"seed_demo refers to a missing file: {rel}")
    return problems


def main() -> int:
    written = build()
    print(f"{'file':<58} {'bytes':>8}  sha256")
    for p in written:
        print(f"{str(p.relative_to(ROOT)):<58} {p.stat().st_size:>8}  {hashlib.sha256(p.read_bytes()).hexdigest()[:16]}")
    problems = verify(written)
    for pr in problems:
        print("PROBLEM:", pr)
    print(f"{len(written)} files written to {OUT.relative_to(ROOT)}; {len(problems)} problems")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
