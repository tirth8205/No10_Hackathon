"""Build the Gateway Check demo deck (.pptx, 16:9) from a JSON manifest.

    python demo/keynote/build_deck.py manifest.json out.pptx --embed
    python demo/keynote/build_deck.py manifest.json out.pptx --placeholders

--embed        each content slide embeds its mp4 (plays on click in PowerPoint/Keynote).
--placeholders each content slide gets a grey box where the video goes, for import
               into Google Slides (which drops embedded video from .pptx); the clips
               are then added by hand with Insert > Video > Google Drive.

Manifest:
  {"title": str, "subtitle": str,
   "slides": [{"heading": str, "section": str (optional, shown in the header band),
               "lines": [str, ...] (optional), "movie": "/abs/path.mp4" or null,
               "caption": str (optional)}],
   "closing": {"heading": str, "lines": [str, ...], "section": str (optional)}}

Needs ffmpeg/ffprobe on PATH (or in /opt/homebrew/bin) for poster frames and clip sizes.
"""

import json
import os
import shutil
import subprocess
import sys
import tempfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

# GOV.UK colours
BLACK = RGBColor(0x0B, 0x0C, 0x0C)
GREY = RGBColor(0x50, 0x5A, 0x5F)
LIGHT_GREY = RGBColor(0xF3, 0xF2, 0xF1)
BLUE = RGBColor(0x1D, 0x70, 0xB8)
YELLOW = RGBColor(0xFF, 0xDD, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"  # GOV.UK's fallback; GDS Transport is not available

SERVICE_NAME = "Gateway Check"
FOOTER_TEXT = "Beta — This is a new service"

# The slide is laid out on a 1920x1080 pixel grid: 1px = 6350 EMU (13.333in x 7.5in).
SLIDE_W, SLIDE_H = 1920, 1080
MARGIN = 60
CONTENT_W = SLIDE_W - 2 * MARGIN
TOP_LINE_H = 10
BAND_H = 90
HEADING_TOP = 140
HEADING_H = 115
MEDIA_TOP = 270
CAPTION_H = 56
FOOTER_TOP = 1015
FOOTER_H = 36
MEDIA_BOTTOM = FOOTER_TOP - 20 - CAPTION_H - 12
HEADING_CHARS_PER_LINE = 36  # roughly what fits at 48pt bold across CONTENT_W
DEFAULT_ASPECT = 16 / 10


def px(n):
    return Emu(int(round(n * 6350)))


def tool(name):
    return shutil.which(name) or f"/opt/homebrew/bin/{name}"


def clip_aspect(path):
    """Width/height of the clip, or DEFAULT_ASPECT if it can't be read."""
    if not path or not os.path.exists(path):
        return DEFAULT_ASPECT
    try:
        out = subprocess.run(
            [tool("ffprobe"), "-v", "error", "-select_streams", "v:0",
             "-show_entries", "stream=width,height", "-of", "csv=p=0", path],
            capture_output=True, text=True, check=True).stdout
        w, h = (int(v) for v in out.strip().split(",")[:2])
        return w / h
    except (OSError, subprocess.CalledProcessError, ValueError):
        return DEFAULT_ASPECT


def poster_frame(path, tmpdir):
    """First frame of the clip as a PNG, used as the still shown before playback."""
    poster = os.path.join(tmpdir, os.path.basename(path) + ".png")
    subprocess.run(
        [tool("ffmpeg"), "-y", "-loglevel", "error", "-i", path,
         "-frames:v", "1", "-update", "1", poster],
        check=True)
    return poster


def heading_block(slide, heading):
    """Draw the 48pt heading; return the y where the content below it starts.
    A heading long enough to wrap gets a second line's worth of room."""
    extra = 110 if len(heading) > HEADING_CHARS_PER_LINE else 0
    text(slide, MARGIN, HEADING_TOP, CONTENT_W, HEADING_H + extra, [heading], 48, BLACK, bold=True)
    return MEDIA_TOP + extra


def media_box(aspect, top):
    """Left, top, width, height (px) of the video area below the heading."""
    h = MEDIA_BOTTOM - top
    w = int(h * aspect)
    if w > CONTENT_W:
        w = CONTENT_W
        h = int(w / aspect)
    return MARGIN, top, w, h


def rect(slide, left, top, width, height, colour):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(left), px(top), px(width), px(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = colour
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def style_run(run, size, colour, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = colour


def text(slide, left, top, width, height, lines, size, colour, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, bullets=False, frame=None):
    """A text box with one paragraph per line. `frame` reuses a shape's own text frame."""
    if frame is None:
        box = slide.shapes.add_textbox(px(left), px(top), px(width), px(height))
        frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for i, line in enumerate(lines):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        style_run(p.add_run(), size, colour, bold)
        p.runs[0].text = line
        if bullets:
            p.space_after = Pt(size * 0.5)
            ppr = p._p.get_or_add_pPr()
            ppr.set("marL", str(px(size * 2.2)))
            ppr.set("indent", str(-px(size * 2.2)))
            bu = ppr.makeelement(qn("a:buChar"), {"char": "•"})
            ppr.append(bu)
    return frame


def header(slide, section):
    rect(slide, 0, 0, SLIDE_W, TOP_LINE_H, BLACK)
    rect(slide, 0, TOP_LINE_H, SLIDE_W, BAND_H, BLUE)
    text(slide, MARGIN, TOP_LINE_H, 900, BAND_H, [SERVICE_NAME], 30, WHITE, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    if section:
        text(slide, SLIDE_W - MARGIN - 900, TOP_LINE_H, 900, BAND_H, [section], 24, WHITE,
             align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def footer(slide, number):
    text(slide, MARGIN, FOOTER_TOP, 1200, FOOTER_H, [FOOTER_TEXT], 18, GREY,
         anchor=MSO_ANCHOR.BOTTOM)
    text(slide, SLIDE_W - MARGIN - 300, FOOTER_TOP, 300, FOOTER_H, [str(number)], 18, GREY,
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.BOTTOM)


def title_slide(prs, manifest):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    header(slide, manifest.get("section", ""))
    text(slide, MARGIN, 140, CONTENT_W, 340, [manifest["title"]], 80, BLACK, bold=True,
         anchor=MSO_ANCHOR.BOTTOM)
    rect(slide, MARGIN, 500, 240, 8, YELLOW)
    if manifest.get("subtitle"):
        text(slide, MARGIN, 540, CONTENT_W, 300, [manifest["subtitle"]], 28, GREY)


def content_slide(prs, spec, number, mode, tmpdir):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, spec.get("section", ""))
    footer(slide, number)
    body_top = heading_block(slide, spec["heading"])
    lines = spec.get("lines") or []
    movie = spec.get("movie")

    if not movie:
        if lines:
            text(slide, MARGIN, body_top + 15, CONTENT_W, MEDIA_BOTTOM - body_top, lines,
                 28, BLACK, bullets=True)
        return

    left, top, width, height = media_box(clip_aspect(movie), body_top)
    if mode == "embed":
        slide.shapes.add_movie(movie, px(left), px(top), px(width), px(height),
                               poster_frame_image=poster_frame(movie, tmpdir),
                               mime_type="video/mp4")
    else:
        box = rect(slide, left, top, width, height, LIGHT_GREY)
        label = f"Video: {os.path.basename(movie)} — Insert > Video > Google Drive"
        text(slide, 0, 0, 0, 0, [label], 24, GREY, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, frame=box.text_frame)
    if spec.get("caption"):
        text(slide, left, top + height + 12, CONTENT_W, CAPTION_H, [spec["caption"]], 24, GREY)
    if lines:
        side_left = left + width + 50
        text(slide, side_left, top, SLIDE_W - MARGIN - side_left, height, lines, 28, BLACK,
             bullets=True)


def closing_slide(prs, spec, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, spec.get("section", ""))
    footer(slide, number)
    body_top = heading_block(slide, spec["heading"])
    if spec.get("lines"):
        text(slide, MARGIN, body_top + 15, CONTENT_W, MEDIA_BOTTOM - body_top, spec["lines"],
             28, BLACK, bullets=True)


def build(manifest, out_path, mode):
    prs = Presentation()
    prs.slide_width = px(SLIDE_W)
    prs.slide_height = px(SLIDE_H)
    with tempfile.TemporaryDirectory() as tmpdir:
        title_slide(prs, manifest)
        number = 1
        for spec in manifest.get("slides", []):
            number += 1
            content_slide(prs, spec, number, mode, tmpdir)
        if manifest.get("closing"):
            closing_slide(prs, manifest["closing"], number + 1)
        prs.save(out_path)


def main(argv):
    flags = {a for a in argv if a.startswith("--")}
    args = [a for a in argv if not a.startswith("--")]
    if len(args) != 2 or len(flags & {"--embed", "--placeholders"}) != 1:
        sys.exit("usage: build_deck.py manifest.json out.pptx (--embed | --placeholders)")
    with open(args[0]) as f:
        manifest = json.load(f)
    mode = "embed" if "--embed" in flags else "placeholders"
    build(manifest, args[1], mode)
    print(f"wrote {args[1]} ({mode})")


if __name__ == "__main__":
    main(sys.argv[1:])
