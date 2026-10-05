"""Turn an uploaded file into text.

PDF  -> pypdf text; if the text is too thin it is treated as a scan and sent to
        the model as a document block for transcription.
DOCX -> paragraphs and tables.
Image-> sent to the model as an image block for transcription.
TXT/EML -> as is.

The original bytes stay in memory only. Nothing is written to disk.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import PurePath

from pypdf import PdfReader

IMAGE_TYPES = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".gif": "image/gif",
    ".webp": "image/webp",
}
TEXT_TYPES = {".txt", ".eml", ".md"}
MIN_CHARS_PER_PAGE = 50


@dataclass
class Ingested:
    text: str
    kind: str  # pdf | docx | image | text
    pages: int = 1
    needs_transcription: bool = False
    media_type: str | None = None
    data: bytes | None = None  # only kept when the model must read the file itself
    notes: list[str] = field(default_factory=list)


class UnsupportedFile(ValueError):
    pass


def ingest(data: bytes, filename: str) -> Ingested:
    ext = PurePath(filename).suffix.lower()
    if ext == ".pdf":
        return _pdf(data)
    if ext == ".docx":
        return _docx(data)
    if ext in IMAGE_TYPES:
        return Ingested(text="", kind="image", needs_transcription=True, media_type=IMAGE_TYPES[ext], data=data)
    if ext in TEXT_TYPES or not ext:
        return Ingested(text=_decode(data), kind="text")
    raise UnsupportedFile(f"unsupported file type {ext!r}; use PDF, DOCX, JPG, PNG or TXT")


def _decode(data: bytes) -> str:
    for enc in ("utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf(data: bytes) -> Ingested:
    reader = PdfReader(io.BytesIO(data))
    pages = [(p.extract_text() or "") for p in reader.pages]
    text = "\n\n".join(t.strip() for t in pages).strip()
    n = max(len(pages), 1)
    thin = len(text) < MIN_CHARS_PER_PAGE * n
    out = Ingested(text=text, kind="pdf", pages=n)
    if thin:
        out.needs_transcription = True
        out.media_type = "application/pdf"
        out.data = data
        out.notes.append("little or no text layer; treated as a scan")
    return out


def _docx(data: bytes) -> Ingested:
    from docx import Document  # imported lazily; python-docx is slow to import

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return Ingested(text="\n".join(parts), kind="docx")
